"""Read-only input adapters; reference labels and visit metadata stay separate."""

import math
import os
from dataclasses import dataclass, field
from numbers import Integral, Real
from pathlib import Path
from typing import Dict, List, Mapping, Optional, Sequence, Tuple, Union


PathInput = Union[str, os.PathLike]


@dataclass
class VisitInput:
    """One visit's inputs; grouping/time metadata are not model features.

    ``image`` may hold original bytes, a path, or an already decoded image.
    Other fields carry audit metadata and never enter the image classifier.
    Targets deliberately have no field in this object.
    """

    sample_id: str
    image: object = None
    structured: Mapping[str, object] = field(default_factory=dict)
    text: Optional[str] = None
    human: Mapping[str, object] = field(default_factory=dict)
    patient_id: Optional[str] = None
    eye_id: Optional[str] = None
    time_years: Optional[float] = None


@dataclass
class GrapeDataset:
    """All GRAPE visits and separate eye-level progression reference labels.

    The labels are progression outcomes, never glaucoma/normal diagnoses.
    The classifier uses only raw CFPs. Contemporaneous IOP is retained for
    availability accounting; VF measurements remain separate reference outcomes.
    """

    visits: List[VisitInput]
    progression_labels: Dict[str, Dict[str, int]]

    def image_visits(self, min_visits: int = 3) -> List[VisitInput]:
        """Select eyes with at least ``min_visits`` available original CFPs.

        This is a cohort eligibility rule, not a predictive feature. Missing
        image visits are omitted; no follow-up visit is duplicated as baseline.
        """
        if isinstance(min_visits, bool) or not isinstance(min_visits, Integral) or min_visits < 1:
            raise ValueError("min_visits must be a positive integer")
        counts: Dict[str, int] = {}
        for visit in self.visits:
            if visit.image is not None:
                counts[visit.eye_id] = counts.get(visit.eye_id, 0) + 1
        return sorted(
            (v for v in self.visits if v.image is not None and counts[v.eye_id] >= min_visits),
            key=_visit_sort_key,
        )


def _visit_sort_key(visit: VisitInput):
    return (visit.patient_id or "", visit.eye_id or "", visit.time_years, visit.sample_id)


def _binary_label(value: object, context: str) -> int:
    if isinstance(value, bool) or not isinstance(value, Integral) or value not in (0, 1):
        raise ValueError(f"{context}: label must be an integer 0 or 1, got {value!r}")
    return int(value)


def _sample_id(value: object, seen: set, context: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{context}: sample_id must be a nonempty string")
    if value in seen:
        raise ValueError(f"{context}: duplicate sample_id {value!r}")
    seen.add(value)
    return value


def _grape_rows(sheet, required: Sequence[str]) -> List[Dict[str, object]]:
    """Flatten GRAPE's two header rows without depending on pandas inference."""
    rows = iter(sheet.iter_rows(values_only=True))
    first = next(rows, None)
    second = next(rows, None)
    if first is None or second is None:
        raise ValueError(f"{sheet.title}: expected two header rows")
    names = []
    parent = None
    for top, child in zip(first, second):
        if top is not None:
            parent = str(top).strip()
        name = parent
        if child is not None:
            name = f"{parent}/{child}"
        names.append(name)
    missing = set(required) - set(names)
    if missing:
        raise ValueError(f"{sheet.title}: missing GRAPE columns {sorted(missing)}")
    for name in required:
        if names.count(name) != 1:
            raise ValueError(f"{sheet.title}: ambiguous duplicate column {name!r}")
    return [dict(zip(names, row)) for row in rows if any(value is not None for value in row)]


def _grape_eye(row: Mapping[str, object], context: str) -> Tuple[str, str]:
    patient = row["Subject Number"]
    if isinstance(patient, Integral) and not isinstance(patient, bool):
        patient = str(patient)
    elif isinstance(patient, str) and patient.strip():
        patient = patient.strip()
    else:
        raise ValueError(f"{context}: Subject Number must be an integer or nonempty string")
    laterality = row["Laterality"]
    if laterality not in ("OD", "OS"):
        raise ValueError(f"{context}: Laterality must be OD or OS")
    return patient, f"{patient}_{laterality}"


def load_grape(root: PathInput = "data/raw/grape") -> GrapeDataset:
    """Read GRAPE's retained Excel workbook and original CFP references.

    ``Follow-up`` already includes baseline visits, so only that sheet creates
    VisitInput records. ``Baseline`` supplies PLR2, PLR3 and MD-slope progression
    labels separately. The workbook's ``/`` image marker becomes ``None``.
    Available filenames must resolve to an existing original ``extracted/CFPs``
    file. Annotation overlays, VF, total visit counts and baseline-only RNFL
    measurements are never copied into model feature channels.
    """
    import openpyxl

    root = Path(root).resolve()
    workbook_path = root / "files" / "VF and clinical information.xlsx"
    image_directory = (root / "extracted" / "CFPs").resolve()
    outcome_columns = {
        "plr2": "Progression Status/PLR2",
        "plr3": "Progression Status/PLR3",
        "md_slope": "Progression Status/MD",
    }
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if not {"Baseline", "Follow-up"}.issubset(workbook.sheetnames):
            raise ValueError("GRAPE workbook requires Baseline and Follow-up sheets")
        baseline = _grape_rows(
            workbook["Baseline"], ["Subject Number", "Laterality"] + list(outcome_columns.values())
        )
        followup = _grape_rows(
            workbook["Follow-up"],
            ["Subject Number", "Laterality", "Visit Number", "Interval Years", "IOP", "Corresponding CFP"],
        )
    finally:
        workbook.close()

    outcomes: Dict[str, Dict[str, int]] = {}
    for row_number, row in enumerate(baseline, start=3):
        context = f"Baseline row {row_number}"
        _, eye_id = _grape_eye(row, context)
        if eye_id in outcomes:
            raise ValueError(f"{context}: duplicate baseline eye_id {eye_id!r}")
        outcomes[eye_id] = {
            name: _binary_label(row[column], f"{context}, {column}")
            for name, column in outcome_columns.items()
        }

    visits: List[VisitInput] = []
    seen: set = set()
    visit_order: Dict[str, List[Tuple[int, float]]] = {}
    for row_number, row in enumerate(followup, start=3):
        context = f"Follow-up row {row_number}"
        patient_id, eye_id = _grape_eye(row, context)
        if eye_id not in outcomes:
            raise ValueError(f"{context}: eye_id {eye_id!r} has no baseline progression labels")
        number = row["Visit Number"]
        if isinstance(number, bool) or not isinstance(number, Integral) or number < 1:
            raise ValueError(f"{context}: Visit Number must be a positive integer")
        sample_id = _sample_id(f"{eye_id}_{number}", seen, context)
        elapsed = row["Interval Years"]
        if isinstance(elapsed, bool) or not isinstance(elapsed, Real) or not math.isfinite(elapsed) or elapsed < 0:
            raise ValueError(f"{context}: Interval Years must be a finite nonnegative number")
        elapsed = float(elapsed)
        iop = row["IOP"]
        if iop is not None:
            if isinstance(iop, bool) or not isinstance(iop, Real) or not math.isfinite(iop):
                raise ValueError(f"{context}: IOP must be a finite number or null")
            iop = float(iop)
        filename = row["Corresponding CFP"]
        image = None
        if filename not in (None, "", "/"):
            if not isinstance(filename, str) or Path(filename).name != filename:
                raise ValueError(f"{context}: Corresponding CFP must be an original CFP filename")
            image = (image_directory / filename).resolve()
            if image.parent != image_directory:
                raise ValueError(f"{context}: Corresponding CFP must stay inside extracted/CFPs")
            if not image.is_file():
                raise FileNotFoundError(f"{context}: original CFP is missing: {image}")
        visits.append(VisitInput(
            sample_id=sample_id, image=image, structured={"iop": iop},
            patient_id=patient_id, eye_id=eye_id, time_years=elapsed,
        ))
        visit_order.setdefault(eye_id, []).append((int(number), elapsed))
    if not visits:
        raise ValueError("GRAPE workbook has no follow-up visit records")
    missing_visits = set(outcomes) - set(visit_order)
    if missing_visits:
        raise ValueError(f"Baseline eyes have no follow-up records: {sorted(missing_visits)}")
    for eye_id, entries in visit_order.items():
        entries.sort()
        if any(right[1] <= left[1] for left, right in zip(entries, entries[1:])):
            raise ValueError(f"{eye_id}: Interval Years must increase with Visit Number")
    return GrapeDataset(visits=sorted(visits, key=_visit_sort_key), progression_labels=outcomes)
