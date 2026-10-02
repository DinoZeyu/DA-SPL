"""Small synthetic fixtures only; these tests never load the research cohort."""

import tempfile
import unittest
from pathlib import Path

import openpyxl

from glaboost.data import (
    VisitInput,
    load_grape,
)


class DataAdapterTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.directory = Path(self.temporary.name)


    def grape(self, baseline=None, followup=None):
        if baseline is None:
            baseline = [
                [1, "OD", 1, 0, 1, "OAG", 4, 75, 18],
                [1, "OS", 0, 0, 0, "OAG", 2, 80, 22],
            ]
        if followup is None:
            # Deliberately unordered, with one unavailable image and two eyes.
            followup = [
                [1, "OD", 3, 2.0, 16, "1_OD_3.jpg", 17],
                [1, "OS", 2, 1.5, 14, "1_OS_2.jpg", 22],
                [1, "OD", 1, 0.0, 15, "1_OD_1.jpg", 18],
                [1, "OD", 4, 3.0, 17, "/", 16],
                [1, "OS", 1, 0.0, 13, "1_OS_1.jpg", 23],
                [1, "OD", 2, 0.5, 16, "1_OD_2.jpg", 18],
            ]
        root = self.directory / "grape"
        (root / "files").mkdir(parents=True, exist_ok=True)
        (root / "extracted" / "CFPs").mkdir(parents=True, exist_ok=True)
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Baseline"
        sheet.append([
            "Subject Number", "Laterality", "Progression Status", None, None,
            "Category of Glaucoma", "Total Visits", "OCT RNFL thickness", "VF",
        ])
        sheet.append([None, None, "PLR2", "PLR3", "MD", None, None, "Mean", 0])
        for row in baseline:
            sheet.append(row)
        sheet = book.create_sheet("Follow-up")
        sheet.append([
            "Subject Number", "Laterality", "Visit Number", "Interval Years",
            "IOP", "Corresponding CFP", "VF",
        ])
        sheet.append([None, None, None, None, None, None, 0])
        for row in followup:
            sheet.append(row)
            if row[5] not in (None, "", "/"):
                # Empty placeholders prove that loading references never decodes images.
                (root / "extracted" / "CFPs" / row[5]).touch()
        book.save(root / "files" / "VF and clinical information.xlsx")
        book.close()
        return root

    def test_grape_uses_original_cfps_and_does_not_duplicate_baseline(self):
        root = self.grape()
        workbook = root / "files" / "VF and clinical information.xlsx"
        before = workbook.read_bytes()
        dataset = load_grape(root)
        self.assertEqual(len(dataset.visits), 6)
        self.assertEqual(dataset.visits[0].image, (root / "extracted/CFPs/1_OD_1.jpg").resolve())
        missing = next(v for v in dataset.visits if v.sample_id == "1_OD_4")
        self.assertIsNone(missing.image)
        self.assertEqual(workbook.read_bytes(), before)

    def test_grape_separates_outcomes_metadata_and_iop_features(self):
        dataset = load_grape(self.grape())
        self.assertEqual(dataset.progression_labels, {
            "1_OD": {"plr2": 1, "plr3": 0, "md_slope": 1},
            "1_OS": {"plr2": 0, "plr3": 0, "md_slope": 0},
        })
        self.assertEqual({v.patient_id for v in dataset.visits}, {"1"})
        self.assertEqual({v.eye_id for v in dataset.visits}, {"1_OD", "1_OS"})
        for visit in dataset.visits:
            self.assertEqual(set(visit.structured), {"iop"})
            self.assertIsNone(visit.text)
            self.assertEqual(visit.human, {})
            self.assertFalse(hasattr(visit, "target"))
        self.assertEqual([v.time_years for v in dataset.visits[:4]], [0, 0.5, 2, 3])

    def test_grape_image_eligibility_counts_only_available_images(self):
        dataset = load_grape(self.grape())
        self.assertEqual([v.sample_id for v in dataset.image_visits()], ["1_OD_1", "1_OD_2", "1_OD_3"])
        self.assertEqual(len(dataset.image_visits(min_visits=2)), 5)
        self.assertEqual(dataset.image_visits(min_visits=4), [])
        for invalid in (0, -1, True, 2.5):
            with self.subTest(min_visits=invalid):
                with self.assertRaises(ValueError):
                    dataset.image_visits(min_visits=invalid)

    def test_grape_rejects_missing_referenced_original_image(self):
        root = self.grape()
        (root / "extracted/CFPs/1_OD_1.jpg").unlink()
        with self.assertRaisesRegex(FileNotFoundError, "original CFP is missing"):
            load_grape(root)

    def test_grape_rejects_invalid_progression_labels(self):
        for invalid in (2, None, True):
            with self.subTest(label=invalid):
                baseline = [[1, "OD", invalid, 0, 1, "OAG", 4, 75, 18]]
                with self.assertRaisesRegex(ValueError, "label"):
                    load_grape(self.grape(baseline=baseline))

    def test_grape_requires_matching_unique_baseline_eyes(self):
        one = [1, "OD", 1, 0, 1, "OAG", 4, 75, 18]
        for baseline, message in [
            ([one], "no baseline progression labels"),
            ([one, one], "duplicate baseline eye_id"),
        ]:
            with self.subTest(baseline=baseline):
                with self.assertRaisesRegex(ValueError, message):
                    load_grape(self.grape(baseline=baseline))

    def test_grape_rejects_repeated_visits_or_reversed_time(self):
        baseline = [[1, "OD", 1, 0, 1, "OAG", 2, 75, 18]]
        for followup, message in [
            ([[1, "OD", 1, 0, 15, "/", 18], [1, "OD", 1, 1, 15, "/", 18]], "duplicate sample_id"),
            ([[1, "OD", 1, 1, 15, "/", 18], [1, "OD", 2, 0, 15, "/", 18]], "must increase"),
        ]:
            with self.subTest(followup=followup):
                with self.assertRaisesRegex(ValueError, message):
                    load_grape(self.grape(baseline=baseline, followup=followup))

    def test_grape_rejects_baseline_without_any_visit(self):
        with self.assertRaisesRegex(ValueError, "no follow-up records"):
            load_grape(self.grape(followup=[[1, "OD", 1, 0, 15, "/", 18]]))


    def test_visit_mapping_defaults_are_independent(self):
        first, second = VisitInput("a"), VisitInput("b")
        first.structured["a"] = 1
        first.human["b"] = 2
        self.assertEqual(second.structured, {})
        self.assertEqual(second.human, {})


if __name__ == "__main__":
    unittest.main()
